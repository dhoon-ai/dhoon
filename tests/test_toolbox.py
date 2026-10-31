"""Unit tests for the toolbox. They need no internet, Docker or Ollama.

Run:  pip install -r toolbox/requirements.txt -r requirements-dev.txt && pytest
"""

import importlib
import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "toolbox"))


@pytest.fixture()
def tb(tmp_path, monkeypatch):
    """Import toolbox_server with an isolated temporary workspace."""
    monkeypatch.setenv("WORKSPACE", str(tmp_path / "ws"))
    monkeypatch.setenv("FILES_KEY", str(tmp_path / "files.key"))  # not the real key
    import toolbox_server

    return importlib.reload(toolbox_server)


# ---------------------------------------------------------------- general ---
def test_current_datetime_has_year(tb):
    assert "20" in tb.current_datetime()


def test_all_tools_registered(tb):
    import asyncio

    names = {t.name for t in asyncio.run(tb.mcp.list_tools())}
    assert {
        "current_datetime",
        "web_search",
        "read_webpage",
        "list_files",
        "read_file",
        "write_file",
        "run_python",
        "stock_price_summary",
        "stock_indicators",
        "stock_fundamentals",
    } <= names


# ---------------------------------------------------------------- files ---
def test_write_read_list(tb):
    assert "Saved" in tb.write_file("notes/a.md", "hello")
    assert tb.read_file("notes/a.md") == "hello"
    assert "notes/a.md" in tb.list_files("notes")


@pytest.mark.parametrize("bad", ["../secret.txt", "/etc/passwd", "notes/../../x"])
def test_paths_outside_workspace_are_blocked(tb, bad):
    with pytest.raises(ValueError):
        tb.write_file(bad, "x")


def test_read_file_truncates(tb):
    tb.write_file("big.txt", "x" * (tb.MAX_TEXT + 100))
    assert tb.read_file("big.txt").endswith("[truncated]")


# ---------------------------------------------------------------- web ---
class FakeResponse:
    def __init__(self, text="", data=None):
        self.text, self._data = text, data

    def raise_for_status(self):
        pass

    def json(self):
        return self._data


def test_web_search_formats_results(tb, monkeypatch):
    data = {"results": [{"title": f"T{i}", "url": f"https://e.com/{i}", "content": "c" * 500} for i in range(12)]}
    monkeypatch.setattr(tb.httpx, "get", lambda *a, **k: FakeResponse(data=data))
    out = json.loads(tb.web_search("anything", max_results=3))
    assert len(out) == 3 and out[0]["url"] == "https://e.com/0"
    assert len(out[0]["snippet"]) == 300


def test_web_search_empty(tb, monkeypatch):
    monkeypatch.setattr(tb.httpx, "get", lambda *a, **k: FakeResponse(data={"results": []}))
    assert tb.web_search("nothing") == "No results."


def test_read_webpage_strips_scripts_and_marks_untrusted(tb, monkeypatch):
    html = b"<html><script>evil()</script><nav>menu</nav><p>Real text</p></html>"
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("text/html; charset=utf-8", html))
    out = tb.read_webpage("https://example.com")
    assert out.startswith("[UNTRUSTED WEB CONTENT]")
    assert "Real text" in out and "evil" not in out and "menu" not in out


def test_read_webpage_reads_online_pdf_and_text(tb, monkeypatch):
    import io

    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(40, 800, "Quarterly report: revenue grew 12 percent.")
    c.showPage()
    c.save()
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("application/pdf", buf.getvalue()))
    assert "revenue grew 12 percent" in tb.read_webpage("https://example.com/r.pdf")
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("text/plain", b"plain words here"))
    assert "plain words here" in tb.read_webpage("https://example.com/a.txt")
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("image/png", b"\x89PNG"))
    assert "not a text page" in tb.read_webpage("https://example.com/a.png")


@pytest.mark.parametrize(
    "url,ip",
    [("http://localhost:3080", "127.0.0.1"), ("http://router.home/", "192.168.1.1"), ("http://x.test/", "10.0.0.5")],
)
def test_read_webpage_blocks_local_network(tb, monkeypatch, url, ip):
    monkeypatch.setattr(tb.socket, "getaddrinfo", lambda host, port: [(2, 1, 6, "", (ip, 0))])
    with pytest.raises(ValueError, match="home network"):
        tb._check_public_url(url)


def test_read_webpage_rejects_non_web_schemes(tb):
    with pytest.raises(ValueError, match="http"):
        tb._check_public_url("file:///etc/passwd")


def test_mcp_returns_explanation_instead_of_generic_error(tb, monkeypatch):
    """The model must see WHY a tool failed, not just 'Error executing tool'."""
    import asyncio

    def refuse(url, headers=None):
        req = tb.httpx.Request("GET", url)
        raise tb.httpx.HTTPStatusError("403", request=req, response=tb.httpx.Response(403, request=req))

    monkeypatch.setattr(tb, "_fetch", refuse)
    result = asyncio.run(tb.mcp.call_tool("read_webpage", {"url": "https://example.com"}))
    text = str(result)
    assert "refuses automated readers" in text and "HTTP 403" in text
    assert "honest identity" in text and "different search result" in text and "tell the user" in text
    result = asyncio.run(tb.mcp.call_tool("write_file", {"path": "../../etc/x", "content": "x"}))
    assert "outside the workspace" in str(result)


# ---------------------------------------------------------------- sandbox ---
def test_run_python_uses_locked_down_docker(tb, monkeypatch):
    seen = {}

    def fake_run(cmd, **kw):
        seen["cmd"], seen["input"] = cmd, kw["input"]
        return subprocess.CompletedProcess(cmd, 0, stdout="42\n", stderr="")

    monkeypatch.setattr(tb.subprocess, "run", fake_run)
    assert tb.run_python("print(6*7)") == "42"
    cmd = " ".join(seen["cmd"])
    for flag in ("--network none", "--read-only", "--cap-drop ALL", "no-new-privileges"):
        assert flag in cmd
    assert seen["input"].endswith("print(6*7)") and "import pandas as pd" in seen["input"]


def test_run_python_timeout(tb, monkeypatch):
    def boom(cmd, **kw):
        raise subprocess.TimeoutExpired(cmd, 1)

    monkeypatch.setattr(tb.subprocess, "run", boom)
    assert "too long" in tb.run_python("while True: pass")


# ---------------------------------------------------------------- stocks ---
def _fake_prices(days=260, start=100.0, step=0.5):
    idx = pd.date_range("2025-01-01", periods=days, freq="B")
    close = [start + i * step for i in range(days)]
    return pd.DataFrame({"Close": close, "Volume": [1000] * days}, index=idx)


def test_stock_price_summary(tb, monkeypatch):
    monkeypatch.setattr(tb, "_history", lambda t, p: _fake_prices(days=11, step=1.0))
    out = json.loads(tb.stock_price_summary("TEST", "1mo"))
    assert out["first_close"] == 100 and out["last_close"] == 110
    assert out["period_return_pct"] == 10.0
    assert out["max_drawdown_pct"] == 0.0


def test_stock_indicators_uptrend(tb, monkeypatch):
    monkeypatch.setattr(tb, "_history", lambda t, p: _fake_prices())
    out = json.loads(tb.stock_indicators("TEST"))
    assert out["sma50"] > out["sma200"]
    assert "uptrend" in out["trend"]
    assert out["rsi14"] == 100.0


def test_stock_indicators_not_enough_data(tb, monkeypatch):
    monkeypatch.setattr(tb, "_history", lambda t, p: _fake_prices(days=50))
    assert "Not enough history" in tb.stock_indicators("TEST")


def test_stock_unknown_symbol(tb, monkeypatch):
    monkeypatch.setattr(tb, "_history", lambda t, p: pd.DataFrame())
    assert "No data" in tb.stock_price_summary("NOPE")


# ---------------------------------------------------------------- planner ---
class FakePost:
    def __init__(self, data):
        self._data = data

    def raise_for_status(self):
        pass

    def json(self):
        return self._data


def _ollama(content):
    return FakePost({"message": {"content": content}})


def test_ask_planner_uses_local_instruct_model_once(tb, monkeypatch):
    seen = []

    def fake_post(url, json=None, timeout=None):
        seen.append(json)
        return _ollama("<think>x</think>Checklist:\n▶ Add\nNext steps:\n1. run_python: 2+2")

    monkeypatch.setattr(tb.httpx, "post", fake_post)
    monkeypatch.setattr(tb.httpx, "get", lambda *a, **k: pytest.fail("no web search unless asked"))
    out = tb.ask_planner("add numbers", context="none")
    assert len(seen) == 1 and seen[0]["think"] is False and seen[0]["stream"] is False
    assert seen[0]["model"] == tb.PLANNER_MODEL and "thinking" not in tb.PLANNER_MODEL
    assert "keep_alive" not in seen[0]  # server setting decides (a value here would un-pin the model)
    assert out.startswith("NEXT STEPS FROM PLANNER") and "1. run_python" in out and "<think>" not in out
    assert "call ask_planner again" in out and "progress" in out
    system = seen[0]["messages"][0]["content"]
    assert tb.dt.date.today().isoformat() in system  # planner knows today's date
    assert "current data" in system and "Next steps" in system
    assert "Checklist:" in system and "✅" in system and "▶" in system and "⬜" in system
    assert "Progress:" in out and "show the user the Checklist" in out


def test_ask_planner_second_round_sends_progress_and_detects_finished(tb, monkeypatch):
    seen = {}

    def fake_post(url, json=None, timeout=None):
        seen["user"] = json["messages"][1]["content"]
        return _ollama("FINISHED: a table comparing the three tools")

    monkeypatch.setattr(tb.httpx, "post", fake_post)
    monkeypatch.setattr(tb.httpx, "get", lambda *a, **k: pytest.fail("research is for round 1 only"))
    out = tb.ask_planner("compare backup tools", progress="found restic, borg, duplicati", research_query="how to x")
    assert "Task: compare backup tools" in seen["user"] and "found restic, borg" in seen["user"]
    assert "ready to answer" in out and "FINISHED: a table" in out and "write the final answer" in out
    assert "show the user the Checklist" in out  # final progress line before the answer


def test_ask_planner_ollama_down_or_empty(tb, monkeypatch):
    def down(*a, **k):
        raise tb.httpx.ConnectError("down")

    monkeypatch.setattr(tb.httpx, "post", down)
    assert "Planner unavailable" in tb.ask_planner("x")
    monkeypatch.setattr(tb.httpx, "post", lambda *a, **k: _ollama(""))
    assert "Make a simple plan yourself" in tb.ask_planner("x")


@pytest.mark.parametrize(
    "text,secret",
    [
        ("email rahul.k@example.com my report", "rahul.k@example.com"),
        ("see https://bank.example.com/acct?id=99 now", "bank.example.com"),
        ("server 192.168.1.20 backup", "192.168.1.20"),
        ("backup /home/rahul/photos/2024 to disk", "/home/rahul"),
        ("summarise salary_2025.xlsx for me", "salary_2025"),
        ("call +1 (555) 123-4567 tomorrow", "555"),
        ("budget of $12,500 for a car", "12,500"),
        ("account 123456789 balance", "123456789"),
    ],
)
def test_redact_removes_personal_details(tb, text, secret):
    out = tb._redact(text)
    assert secret not in out and "[" in out


def test_redact_keeps_generic_questions_and_removes_user_name(tb, monkeypatch):
    q = "how to compare open source backup tools for linux in 2026"
    assert tb._redact(q) == q
    monkeypatch.setattr(tb.Path, "home", lambda: tb.Path("/home/rahulk"))
    assert "rahulk" not in tb._redact("plan for rahulk backups")
    assert len(tb._redact("word " * 100)) <= 150


def test_ask_planner_web_research_is_redacted_and_reported(tb, monkeypatch):
    searched, planned = [], {}

    def fake_get(url, params=None, timeout=None):
        searched.append(params)
        return FakeResponse(data={"results": [{"title": "Guide", "url": "https://g.example/a", "content": "Step A"}]})

    def fake_post(url, json=None, timeout=None):
        planned["user"] = json["messages"][1]["content"]
        return _ollama("Checklist:\n▶ Research\nNext steps:\n1. web_search: restic")

    monkeypatch.setattr(tb.httpx, "get", fake_get)
    monkeypatch.setattr(tb.httpx, "post", fake_post)
    monkeypatch.setattr(tb, "RESEARCH_ENGINES", "startpage,duckduckgo")
    out = tb.ask_planner("back up my files", research_query="how to back up /home/rahul for me@x.com")
    assert len(searched) == 1 and searched[0]["engines"] == "startpage,duckduckgo"
    assert "/home/rahul" not in searched[0]["q"] and "me@x.com" not in searched[0]["q"]
    assert "Step A" in planned["user"] and "untrusted" in planned["user"]
    assert out.startswith('Searched the web for: "how to back up [path] for [email]"')
    assert "NEXT STEPS FROM PLANNER" in out


def test_research_falls_back_to_all_engines_then_plans_without(tb, monkeypatch):
    searched = []

    def fake_get(url, params=None, timeout=None):
        searched.append(params)
        return FakeResponse(data={"results": []})

    monkeypatch.setattr(tb.httpx, "get", fake_get)
    monkeypatch.setattr(tb.httpx, "post", lambda *a, **k: _ollama("1. web_search: x"))
    monkeypatch.setattr(tb, "RESEARCH_ENGINES", "mojeek")
    out = tb.ask_planner("task", research_query="how to plan a garden")
    assert [p.get("engines") for p in searched] == ["mojeek", None]
    assert "found nothing useful" in out and "1. web_search" in out


# ---------------------------------------------------------------- documents ---
def _bag_of_words(text, dims=64):
    """Stand-in for the embedding model: hashed word counts (similar text -> similar vector)."""
    import hashlib

    vec = [0.0] * dims
    for word in re.findall(r"[a-z]+", text.lower().split(":", 1)[-1]):
        vec[int(hashlib.md5(word.encode()).hexdigest(), 16) % dims] += 1.0
    return vec


@pytest.fixture()
def fake_embed(tb, monkeypatch):
    calls = []

    def fake_post(url, json=None, timeout=None):
        assert url.endswith("/api/embed") and json["model"] == tb.EMBED_MODEL
        calls.append(len(json["input"]))
        return FakePost({"embeddings": [_bag_of_words(t) for t in json["input"]]})

    monkeypatch.setattr(tb.httpx, "post", fake_post)
    return calls


def _make_pdf(path, pages):
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path))
    for text in pages:
        y = 800
        for line in text.split("\n"):
            c.drawString(40, y, line)
            y -= 14
        c.showPage()
    c.save()


def test_search_document_pdf_finds_right_page_and_caches(tb, fake_embed):
    pages = [
        "Alice was beginning to get very tired of sitting by her sister on the bank.",
        "The Cheshire Cat grinned and slowly vanished, leaving only the grin behind.",
        "The Queen of Hearts shouted off with her head at the croquet game.",
    ]
    _make_pdf(tb.WORKSPACE / "alice.pdf", pages)
    out = tb.search_document("alice.pdf", "the cat that vanished with a grin", top_k=1)
    assert "[page 2" in out and "Cheshire" in out
    first_calls = len(fake_embed)
    tb.search_document("alice.pdf", "queen croquet", top_k=1)
    assert len(fake_embed) == first_calls + 1  # index cached: only the query is embedded
    assert ".index" not in tb.list_files(".")  # cache folder is hidden


def test_read_document_pages_and_total(tb):
    _make_pdf(tb.WORKSPACE / "book.pdf", ["first page text", "second page text", "third page text"])
    out = tb.read_document("book.pdf", start_page=2, max_pages=1)
    assert "3 pages in total" in out and "second page" in out and "third page" not in out


def test_search_document_docx_and_txt(tb, fake_embed):
    import docx

    d = docx.Document()
    d.add_paragraph("Invoices are due within thirty days of delivery.")
    d.add_paragraph("Refunds require the original receipt and the product packaging.")
    d.save(tb.WORKSPACE / "policy.docx")
    assert "Refunds" in tb.search_document("policy.docx", "what do I need for a refund receipt", top_k=1)
    (tb.WORKSPACE / "notes.txt").write_text("Backup the server every Sunday night.\n" * 3)
    assert "Sunday" in tb.search_document("notes.txt", "when is the backup", top_k=1)


def test_excel_files_are_read_sheet_by_sheet(tb):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Budget"
    wb.active.append(["Item", "Amount"])
    wb.active.append(["Rent", 1800])
    wb.create_sheet("Notes").append(["Pay rent on the 1st"])
    wb.save(tb.WORKSPACE / "budget.xlsx")
    out = tb.read_document("budget.xlsx")
    assert "2 pages in total" in out and "Sheet: Budget" in out and "Rent\t1800" in out and "Pay rent" in out


def test_sandbox_can_make_excel_and_word_files():
    docker = (Path(__file__).resolve().parents[1] / "sandbox/Dockerfile").read_text()
    assert "openpyxl" in docker and "python-docx" in docker
    import toolbox_server as tbs

    assert "openpyxl" in tbs.run_python.__doc__ and "python-docx" in tbs.run_python.__doc__


def test_document_errors(tb, fake_embed):
    assert "File not found" in tb.search_document("missing.pdf", "x")
    (tb.WORKSPACE / "photo.png").write_bytes(b"\x89PNG")
    assert "Unsupported file type" in tb.search_document("photo.png", "x")
    _make_pdf(tb.WORKSPACE / "blank.pdf", [""])
    assert "No text found" in tb.search_document("blank.pdf", "x")
    with pytest.raises(ValueError):
        tb.read_document("../outside.pdf")


def test_wikipedia_links_use_the_official_api_with_honest_identity(tb, monkeypatch):
    seen = {}

    def fake_fetch(url, headers=None):
        seen["url"], seen["headers"] = url, headers
        body = {"query": {"pages": [{"title": "Rsync", "extract": "rsync is a utility for transferring files."}]}}
        return "application/json", json.dumps(body).encode()

    monkeypatch.setattr(tb, "_fetch", fake_fetch)
    out = tb.read_webpage("https://en.m.wikipedia.org/wiki/Rsync#History")
    assert "rsync is a utility" in out and out.startswith("[UNTRUSTED WEB CONTENT]")
    assert seen["url"].startswith("https://en.wikipedia.org/w/api.php?") and "titles=Rsync" in seen["url"]
    assert "Dhoon" in seen["headers"]["User-Agent"]


def test_wikipedia_missing_article(tb, monkeypatch):
    body = {"query": {"pages": [{"title": "Nope", "missing": True}]}}
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("application/json", json.dumps(body).encode()))
    with pytest.raises(ValueError, match="no article"):
        tb.read_webpage("https://de.wikipedia.org/wiki/Nope")


def test_refused_page_is_retried_once_with_honest_identity(tb, monkeypatch):
    calls = []

    def fetch(url, headers=None):
        calls.append(headers)
        if headers is None:
            req = tb.httpx.Request("GET", url)
            raise tb.httpx.HTTPStatusError("403", request=req, response=tb.httpx.Response(403, request=req))
        return "text/html", b"<p>Welcome, honest bot</p>"

    monkeypatch.setattr(tb, "_fetch", fetch)
    assert "Welcome, honest bot" in tb.read_webpage("https://example.org/page")
    assert len(calls) == 2 and "Dhoon" in calls[1]["User-Agent"]


def test_research_default_uses_searxng_usual_engines_once(tb, monkeypatch):
    searched = []

    def fake_get(url, params=None, timeout=None):
        searched.append(params)
        return FakeResponse(data={"results": []})

    monkeypatch.setattr(tb, "RESEARCH_ENGINES", "")
    monkeypatch.setattr(tb.httpx, "get", fake_get)
    tb._research("how to plan a garden")
    assert len(searched) == 1 and "engines" not in searched[0]  # no engine filter, no duplicate retry


def test_run_python_explains_missing_internet_and_modules(tb, monkeypatch):
    err = "Traceback...\nModuleNotFoundError: No module named 'yfinance'"
    monkeypatch.setattr(tb.subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, "", err))
    out = tb.run_python("import yfinance")
    assert "No module named 'yfinance'" in out and "no internet" in out and "stock_save_history" in out
    doc = tb.run_python.__doc__
    assert "simulated" in doc and "NO internet" in doc


def test_stock_save_history_writes_csv_for_run_python(tb, monkeypatch):
    monkeypatch.setattr(tb, "_history", lambda t, p: pd.DataFrame() if t == "NOPE" else _fake_prices(days=20))
    out = tb.stock_save_history("googl, msft,NOPE", "1mo")
    assert "data/GOOGL_MSFT_1mo.csv" in out and "No data for: NOPE" in out
    df = pd.read_csv(tb.WORKSPACE / "data/GOOGL_MSFT_1mo.csv", parse_dates=["Date"], index_col="Date")
    assert list(df.columns) == ["GOOGL", "MSFT"] and len(df) == 20
    assert "No data" in tb.stock_save_history("NOPE")


def test_every_tool_call_is_logged(tb, monkeypatch, tmp_path):
    import asyncio
    import logging

    logfile = tmp_path / "toolbox.log"
    handler = logging.FileHandler(logfile)
    handler.setFormatter(logging.Formatter("%(message)s"))
    tb.log.addHandler(handler)
    asyncio.run(tb.mcp.call_tool("current_datetime", {}))
    asyncio.run(tb.mcp.call_tool("write_file", {"path": "../../etc/x", "content": "x"}))
    handler.flush()
    text = logfile.read_text()
    assert "CALL  current_datetime" in text and "DONE  current_datetime" in text
    assert "CALL  write_file" in text and "ERROR write_file" in text and "outside the workspace" in text


def test_price_history_drops_rows_without_prices(tb, monkeypatch):
    import sys
    import types

    df = _fake_prices(days=5)
    df.loc[df.index[-1], "Close"] = float("nan")  # Yahoo's empty last row (seen for .NS tickers)

    class FakeTicker:
        def __init__(self, t):
            pass

        def history(self, period, auto_adjust):
            return df.copy()

    monkeypatch.setitem(sys.modules, "yfinance", types.SimpleNamespace(Ticker=FakeTicker))
    out = tb._history("RELIANCE.NS", "1mo")
    assert len(out) == 4 and not out["Close"].isna().any()


ARTICLE_PAGE = (
    "<html><head><title>Bank of Canada cuts rates</title></head><body>"
    "<div class='cookie'>We use cookies to improve your experience. Accept all cookies</div>"
    "<nav><a href='/'>Home</a> <a href='/markets'>Markets</a> <a href='/sports'>Sports</a></nav>"
    "<aside>Most read: Ten celebrity homes you must see. Sign up for our newsletter today!</aside>"
    "<article><h1>Bank of Canada cuts rates</h1>"
    + "".join(
        f"<p>Paragraph {i}: The Bank of Canada lowered its policy rate by a quarter point to 2.25 percent on "
        "Wednesday, citing slowing inflation and a softer job market, and said further cuts depend on data.</p>"
        for i in range(6)
    )
    + "</article><footer>Copyright 2026 Example News. Privacy policy. Terms of use. Contact us.</footer>"
    "<div class='related'>Related: Weather this weekend; Hockey scores; Recipes for fall</div></body></html>"
).encode()


def test_read_webpage_keeps_the_article_and_drops_the_clutter(tb, monkeypatch):
    pytest.importorskip("trafilatura")
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("text/html", ARTICLE_PAGE))
    out = tb.read_webpage("https://news.example/boc")
    assert "lowered its policy rate by a quarter point" in out
    for clutter in ("Accept all cookies", "Sports", "celebrity homes", "Privacy policy", "Recipes for fall"):
        assert clutter not in out
    assert len(out) < len(tb._all_text(ARTICLE_PAGE))


def test_read_webpage_without_an_article_falls_back_to_all_text(tb, monkeypatch):
    page = b"<html><body><ul><li>Item one</li><li>Item two</li></ul></body></html>"
    monkeypatch.setattr(tb, "_fetch", lambda url, headers=None: ("text/html", page))
    assert "Item one" in tb.read_webpage("https://example.org/list")


def test_search_document_also_finds_exact_terms(tb, monkeypatch):
    pytest.importorskip("rank_bm25")
    # an embedding model that can't tell these passages apart: only keyword search can
    monkeypatch.setattr(
        tb.httpx,
        "post",
        lambda url, json=None, timeout=None: FakePost({"embeddings": [[1.0, 0.0] for _ in json["input"]]}),
    )
    notes = "\n\n".join(
        [
            f"Invoice INV-{n} for office chairs was paid in full on time by the finance team."
            for n in range(20000, 20040)
        ]
        + ["Invoice INV-20931 is overdue: the supplier asked for payment by Friday, with a late fee."]
        + [
            f"Invoice INV-{n} for printer paper was paid in full on time by the finance team."
            for n in range(30000, 30040)
        ]
    )
    (tb.WORKSPACE / "invoices.txt").write_text(notes)
    out = tb.search_document("invoices.txt", "INV-20931 overdue", top_k=1)
    assert "INV-20931 is overdue" in out


# ---------------------------------------------------------------- receipts ---
def test_read_receipts_into_a_spreadsheet_and_continue(tb, monkeypatch):
    rr = pytest.importorskip("receipt_reader")
    from PIL import Image
    from receipt_reader import model as rr_model

    folder = tb.WORKSPACE / "receipts"
    folder.mkdir(parents=True)
    for name in ("a.jpg", "b.jpg"):
        Image.new("RGB", (400, 600), "white").save(folder / name)
    good = {"vendor": "Maple Corner Grocery", "date": "2026-10-20", "currency": "CAD",
            "items": [{"description": "Bread", "quantity": 1, "amount": 3.29}], "subtotal": 3.29,
            "taxes": [], "tip": None, "total": 3.29, "payment_method": "cash", "category": "groceries"}
    asked = []
    monkeypatch.setattr(rr_model, "chat", lambda messages, schema, **kw: asked.append(kw) or dict(good))
    monkeypatch.setattr(sys.modules["toolset.receipts"], "MOST", 1)
    first = tb.read_receipts("receipts", "receipts.csv")
    assert "Read 1 receipt(s)" in first and "1 more receipt(s)" in first and asked[0]["model"] == tb.PLANNER_MODEL
    second = tb.read_receipts("receipts", "receipts.csv")
    assert "Read 1 receipt(s)" in second and "more receipt" not in second
    rows = (tb.WORKSPACE / "receipts.csv").read_text().splitlines()
    assert len(rows) == 3 and "Maple Corner Grocery" in rows[2]
    assert "No new receipts" in tb.read_receipts("receipts", "receipts.csv")
    assert rr.__name__ == "receipt_reader"


def test_read_receipts_without_the_package(tb, monkeypatch):
    monkeypatch.setitem(sys.modules, "receipt_reader", None)
    assert "pip install git+https://github.com/knowreva/receipt-reader" in tb.read_receipts()


def test_read_receipts_with_your_usual_currency(tb, monkeypatch):
    pytest.importorskip("receipt_reader")
    from PIL import Image
    from receipt_reader import model as rr_model

    (tb.WORKSPACE / "receipts").mkdir(parents=True)
    Image.new("RGB", (400, 600), "white").save(tb.WORKSPACE / "receipts" / "a.jpg")
    good = {"vendor": "Corner Shop", "date": "2026-10-20", "currency": "$",
            "items": [{"description": "Milk", "quantity": 1, "amount": 5.00}], "subtotal": 5.00,
            "taxes": [], "tip": None, "total": 5.00, "payment_method": "cash", "category": "groceries"}
    prompts = []
    monkeypatch.setattr(rr_model, "chat", lambda messages, schema, **kw: prompts.append(messages[0]["content"]) or dict(good))
    monkeypatch.setenv("RECEIPT_CURRENCY", "cad")
    out = tb.read_receipts("receipts", "r.csv")
    assert "5.0 CAD" in out and "it is CAD" in prompts[0]  # "$" became your usual currency


# ---------------------------------------------------------------- every document at once ---
def test_search_all_documents_finds_the_right_file_and_keeps_up_with_changes(tb, fake_embed):
    import docx

    (tb.WORKSPACE / "contracts").mkdir()
    d = docx.Document()
    d.add_paragraph("This lease renews automatically on March 1 unless either party gives sixty days notice.")
    d.save(tb.WORKSPACE / "contracts" / "lease.docx")
    _make_pdf(tb.WORKSPACE / "contracts" / "phone.pdf", ["The phone plan costs forty dollars a month with no contract."])
    (tb.WORKSPACE / "recipes.txt").write_text("Bake the bread at two hundred degrees for forty minutes.\n" * 3)
    (tb.WORKSPACE / ".hidden.txt").write_text("lease renews secret notes that must not be searched" * 3)
    out = tb.search_all_documents("lease renews automatically", top_k=1)
    assert out.splitlines()[2].startswith("[contracts/lease.docx, page 1") and "March 1" in out
    assert ".hidden" not in out and (tb.WORKSPACE / ".index" / "vectors.sqlite").exists()
    calls = len(fake_embed)
    tb.search_all_documents("bread oven", top_k=1)
    assert len(fake_embed) == calls + 1  # already indexed: only the question is embedded
    assert "recipes.txt" in tb.search_all_documents("bake bread", folder=".", top_k=1)
    assert "recipes.txt" not in tb.search_all_documents("bake bread", folder="contracts", top_k=3)  # one folder
    (tb.WORKSPACE / "recipes.txt").write_text("Soup: simmer the lentils for half an hour.\n" * 3)  # changed
    out = tb.search_all_documents("lentil soup", top_k=1)
    assert "lentils" in out and "bread" not in out
    (tb.WORKSPACE / "contracts" / "phone.pdf").unlink()  # deleted: gone from the index
    assert "phone.pdf" not in tb.search_all_documents("phone plan forty dollars", top_k=5)


def test_new_documents_are_indexed_a_few_at_a_time(tb, fake_embed, monkeypatch):
    from toolset import documents

    monkeypatch.setattr(documents, "NEW_FILES_PER_CALL", 2)
    for i in range(5):
        (tb.WORKSPACE / f"note{i}.txt").write_text(f"Note number {i} about the garden and the tomato plants." * 2)
    out = tb.search_all_documents("tomato garden")
    assert "2 new or changed document(s) indexed for this search; 3 more on the next one" in out
    tb.search_all_documents("tomato garden")
    assert "more on the next one" not in tb.search_all_documents("tomato garden")


def test_the_reranker_reorders_when_switched_on(tb, fake_embed, monkeypatch):
    from toolset import documents

    (tb.WORKSPACE / "a.txt").write_text("Apples are red. " * 40 + "\n\n" + "Pears are green. " * 40)
    plain = tb.search_document("a.txt", "red apples", top_k=1)
    assert "Apples" in plain
    monkeypatch.setattr(documents, "RERANK", True)

    class Stub:  # a cross-encoder that prefers pears
        def rerank(self, query, docs):
            return [5.0 if "Pears" in d else 0.1 for d in docs]

    monkeypatch.setattr(documents, "_reranker", Stub())
    assert "Pears" in tb.search_document("a.txt", "red apples", top_k=1)
    monkeypatch.setattr(documents, "_reranker", None)
    monkeypatch.setitem(sys.modules, "fastembed", None)  # switched on but not installed: the usual order
    assert "Apples" in tb.search_document("a.txt", "red apples", top_k=1)


def test_read_receipts_waits_its_turn_with_long_jobs(tb, monkeypatch):
    """A newsroom run writing with the model while receipts load theirs pushed its model out of memory:
    read_receipts says "not now" while a long job runs, and holds the queue while it reads."""
    pytest.importorskip("receipt_reader")
    import heavy
    from PIL import Image
    from receipt_reader import model as rr_model

    (tb.WORKSPACE / "receipts").mkdir(parents=True)
    Image.new("RGB", (400, 600), "white").save(tb.WORKSPACE / "receipts" / "a.jpg")
    seen = []
    good = {"vendor": "Corner Shop", "date": "2026-10-20", "date_printed": "2026-10-20", "currency": "CAD",
            "items": [{"description": "Milk", "quantity": 1, "amount": 5.00}], "subtotal": 5.00, "taxes": [],
            "tip": None, "total": 5.00, "payment_method": "cash", "category": "groceries"}  # fmt: skip
    monkeypatch.setattr(rr_model, "chat", lambda messages, schema, **kw: seen.append(heavy.running().get("name")) or dict(good))
    real = heavy.running
    monkeypatch.setattr(heavy, "running", lambda: {"name": "newsroom", "since": "10:30", "pid": 1})
    assert tb.read_receipts("receipts", "r.csv").startswith("Not now: newsroom is running (since 10:30)")
    monkeypatch.setattr(heavy, "running", real)  # (not monkeypatch.undo: that would also undo HEAVY_DIR)
    assert "Read 1 receipt(s)" in tb.read_receipts("receipts", "r.csv")
    assert seen == ["receipts"] and heavy.running() == {}  # in the queue while reading, then let go


class _Stream:
    def __init__(self, url, status=200, headers=None, body=b"ok"):
        self.url = url
        self.status_code = status
        self.headers = headers or {}
        self._body = body
        self.is_redirect = status in (301, 302, 303, 307, 308)

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def raise_for_status(self):
        if self.status_code >= 400:
            raise Exception("http error")

    def iter_bytes(self):
        yield self._body


def _addrinfo(host, port):
    ip = "127.0.0.1" if host == "127.0.0.1" else "93.184.216.34"
    return [(2, 1, 6, "", (ip, 0))]


def test_fetch_refuses_redirect_to_private_address(tb, monkeypatch):
    import httpx

    import toolbox.toolset.web as web

    urls = []

    def stream(method, url, **kw):
        urls.append(url)
        if url.endswith("/start"):
            r = _Stream(url, status=302, headers={"location": "http://127.0.0.1/x"}, body=b"")
            r.url = "http://127.0.0.1/x"
            return r
        return _Stream(url, body=b"final")

    monkeypatch.setattr(tb.socket, "getaddrinfo", _addrinfo)
    monkeypatch.setattr(httpx, "stream", stream)
    with pytest.raises(ValueError):
        web._fetch("https://public.example/start")
    assert urls == ["https://public.example/start"]


def test_fetch_follows_relative_redirect(tb, monkeypatch):
    import httpx

    import toolbox.toolset.web as web

    def stream(method, url, **kw):
        if url.endswith("/start"):
            r = _Stream(url, status=302, headers={"location": "/next"}, body=b"")
            r.url = "https://a.example/next"
            return r
        return _Stream(url, body=b"final")

    monkeypatch.setattr(tb.socket, "getaddrinfo", _addrinfo)
    monkeypatch.setattr(httpx, "stream", stream)
    ctype, body = web._fetch("https://a.example/start")
    assert body == b"final"


def test_fetch_too_many_redirects(tb, monkeypatch):
    import httpx

    import toolbox.toolset.web as web

    def stream(method, url, **kw):
        r = _Stream(url, status=302, headers={"location": "/next"}, body=b"")
        return r

    monkeypatch.setattr(tb.socket, "getaddrinfo", _addrinfo)
    monkeypatch.setattr(httpx, "stream", stream)
    with pytest.raises(ValueError, match="too many redirects"):
        web._fetch("https://public.example/start")
